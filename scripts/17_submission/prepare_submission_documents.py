#!/usr/bin/env python3
"""Prepare cleaned SCI submission manuscript and supplementary materials."""

from __future__ import annotations

import re
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[2]
SOURCE_DOCX = PROJECT_ROOT / "manuscript" / "DA-FAP_manuscript.docx"
OUT_DIR = PROJECT_ROOT / "投稿"
SUBMISSION_DOCX = OUT_DIR / "manuscript submission.docx"
SUPPLEMENT_DOCX = OUT_DIR / "supplementary materials.docx"
SUPPLEMENT_XLSX = OUT_DIR / "supplementary source tables.xlsx"


TITLE = (
    "Computational perturbation-guided single-cell, spatial, and multi-cohort "
    "transcriptomics prioritizes degeneration-associated fibro-adipogenic "
    "progenitor programs in skeletal muscle"
)


SUPPLEMENTARY_TABLE_CAPTIONS = {
    1: "Public dataset inventory, metadata fields, inclusion decisions, and rationale.",
    2: "Single-cell and single-nucleus quality-control thresholds and retained cells.",
    3: "Major cell-type marker-module score summaries used for atlas annotation.",
    4: "FAP subtype marker genes ranked by differential expression statistics.",
    5: "DA-FAP and component gene modules used for scoring.",
    6: "Curated ligand-receptor interaction scores across sender and receiver groups.",
    7: "Candidate regulator expression and correlation with DA-FAP component scores.",
    8: "GRN-regression virtual knockdown ranking across candidate regulators.",
    9: "Spatial co-localization statistics in GSE225766 Visium samples.",
    10: "Bulk validation effect sizes across individual cohorts and contrasts.",
    11: "Subtype-level FAP module programs.",
    12: "Leave-one-dataset-out DA-FAP label-recovery sensitivity analysis.",
    13: "DA-FAP module-gene sensitivity and bootstrap stability analysis.",
    14: "Correlations between DA-FAP score and quality-control or artifact-associated covariates.",
    15: "Integration-method sensitivity comparing uncorrected PCA and Harmony representations.",
    16: "Random-effects meta-analysis of module scores across primary human bulk cohorts.",
}


TEXT_REPLACEMENTS = [
    (
        "whether a conserved degeneration-associated FAP program is reproducibly recoverable",
        "whether a recurrent degeneration-associated FAP program is reproducibly recoverable",
    ),
    (
        "These results define a reproducible DA-FAP program",
        "These results prioritize a reproducible DA-FAP program",
    ),
    (
        "a conserved degeneration-associated FAP program is reproducibly recoverable",
        "a recurrent degeneration-associated FAP program is reproducibly recoverable",
    ),
    (
        "testing the hypothesis that human single-cell, spatial, and bulk transcriptomic data contain a conserved degeneration-associated FAP-like program",
        "testing whether human single-cell, spatial, and bulk transcriptomic data support a recurrent degeneration-associated FAP-like program",
    ),
    (
        "the tissue-scale outcome of FAP-state remodeling",
        "a candidate tissue-scale correlate of FAP-state remodeling",
    ),
    (
        "the datasets analysed in the executable pass",
        "the datasets analyzed in the primary analysis",
    ),
    (
        "the executable human single-cell atlas drew on",
        "the primary human single-cell atlas drew on",
    ),
    (
        "The executable human single-cell atlas drew on",
        "The primary human single-cell atlas drew on",
    ),
    (
        "The current executable pass included",
        "The primary analysis included",
    ),
    (
        "current executable pass",
        "primary analysis",
    ),
    (
        "executable pass",
        "primary analysis",
    ),
    (
        "Cross-species spatial analysis localizes DA-FAP programs to dystrophic and injured niches",
        "Cross-species spatial analysis supports DA-FAP localization in dystrophic and injured niches",
    ),
    (
        "Cross-species Visium analysis localized the DA-FAP signature",
        "Cross-species Visium analysis supported localization of the DA-FAP signature",
    ),
    (
        "confirmed higher DA-FAP scores",
        "supported higher DA-FAP scores",
    ),
    (
        "Using a reproducible public-data framework, we defined",
        "Using a reproducible public-data framework, we operationally defined",
    ),
    (
        "a conserved, testable model of stromal remodeling",
        "a recurrent, testable model of stromal remodeling",
    ),
    (
        "represent a reproducible activation state of the muscle stromal compartment detectable wherever degenerative remodeling occurs",
        "represent a recurrent activation state of the muscle stromal compartment detectable across the public degenerative contexts analyzed here",
    ),
    (
        "Positioning DA-FAPs as a druggable state",
        "Positioning DA-FAPs as a therapeutically relevant state",
    ),
    (
        "the core strength of the framework",
        "a major strength of the framework",
    ),
    (
        "A reproducible, public-data analysis identifies a conserved degeneration-associated FAP-like state",
        "A reproducible public-data analysis identifies a recurrent degeneration-associated FAP-like state",
    ),
    (
        "a generalizable disease-associated stromal state",
        "a reproducible disease-associated stromal state",
    ),
    (
        "a prioritized roadmap for experimental and clinical validation",
        "a prioritized framework for experimental and clinical validation",
    ),
    (
        "should be read as hypothesis generation rather than causal proof",
        "should be read as hypothesis generation rather than experimental validation",
    ),
    (
        "The GSE254300 paraspinal cohort, retained as a separate extension because it compares scoliosis-associated paraspinal sides rather than aging or sarcopenia, did not show a positive DA-FAP effect in this first pass.",
        "The GSE254300 paraspinal cohort was retained as a separate extension because it compares scoliosis-associated paraspinal sides rather than aging or sarcopenia; in this analysis, it did not show a positive DA-FAP effect.",
    ),
    (
        "motif-constrained causal frameworks such as CellOracle [10] remain appropriate next-step implementations",
        "motif-constrained causal frameworks such as CellOracle [10] are appropriate follow-up implementations",
    ),
    (
        "All perturbation and regulator results are reported as computational predictions.",
        "All perturbation and regulator results are reported as computational predictions and should be interpreted as hypotheses for experimental testing.",
    ),
    (
        "Processed objects, analysis tables, and figures are available from the corresponding author on reasonable request.",
        "Processed objects, scripts, analysis tables, and figures are organized in the accompanying reproducibility package.",
    ),
    ("Funding. Not applicable.", "Funding. No external funding was received."),
]


SECTION_RENAMES = {
    "Materials and Methods": "Materials and methods",
    "Figures": "Figure legends",
}


def clean_text(text: str) -> str:
    text = text.strip()
    if not text:
        return ""
    if text == "An integrated single-cell, spatial, and multi-cohort bulk transcriptomic atlas with in silico perturbation defines a conserved degeneration-associated fibro-adipogenic progenitor program in skeletal muscle":
        return TITLE
    if text.startswith("Affiliation: ["):
        return ""
    for old, new in TEXT_REPLACEMENTS:
        text = text.replace(old, new)
    text = SECTION_RENAMES.get(text, text)
    text = re.sub(r"\bworkflow\b", "analysis framework", text)
    text = text.replace("first independent human bulk aging validation cohort", "primary independent human bulk aging validation cohort")
    return text


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], font_size: float = 8.5) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:tblHeader")) is None:
                tr_pr.append(OxmlElement("w:tblHeader"))
        for cell_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths_dxa[min(cell_index, len(widths_dxa) - 1)])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                    run.font.size = Pt(font_size)
            if row_index == 0:
                set_cell_shading(cell, "EDEDED")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def apply_base_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size in [("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True


def add_paragraph(doc: Document, text: str, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)
    return paragraph


def add_title_block(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(12)
    run = paragraph.add_run(TITLE)
    run.bold = True
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(16)

    for text in ["Author Name", "Correspondence: author@example.com"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        r.font.size = Pt(12)


def copy_source_table(doc: Document, table, widths: list[int]) -> None:
    out_table = doc.add_table(rows=len(table.rows), cols=len(table.columns))
    out_table.style = "Table Grid"
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            out_table.cell(r, c).text = cell.text
    set_table_geometry(out_table, widths)


def extract_main_parts(source: Document) -> dict[str, list[str]]:
    parts = {"body": [], "table_captions": [], "figure_legends": [], "references": []}
    mode = "body"
    for para in source.paragraphs:
        text = clean_text(para.text)
        if not text:
            continue
        if text == "Tables":
            mode = "table_captions"
            continue
        if text == "Figure legends":
            mode = "figure_legends"
            continue
        if text == "References":
            mode = "references"
            continue
        parts[mode].append(text)
    return parts


def build_submission_manuscript() -> None:
    source = Document(SOURCE_DOCX)
    parts = extract_main_parts(source)
    doc = Document()
    apply_base_styles(doc)
    add_title_block(doc)

    for text in parts["body"][3:]:
        if text in {"Abstract", "Introduction", "Materials and methods", "Results", "Discussion", "Conclusions", "Declarations"}:
            doc.add_paragraph(text, style="Heading 1")
        elif text in {
            "Study design and public data inventory",
            "Single-cell preprocessing, integration, and annotation",
            "FAP subclustering and DA-FAP definition",
            "Trajectory and cell-cell communication",
            "Trajectory and cell–cell communication",
            "Regulatory network and GRN-regression virtual knockdown",
            "Spatial and bulk validation",
            "Robustness and sensitivity analysis",
            "An integrated public atlas resolves the major human skeletal muscle cell populations",
            "FAP-like cells resolve into homeostatic, inflammatory, adipogenic, ECM-remodeling, and DA-FAP states",
            "Pseudotime indicates a modest shift toward DA-FAP states",
            "DA-FAP communication implicates CXCL12-CXCR4 and ECM-integrin axes",
            "DA-FAP communication implicates CXCL12–CXCR4 and ECM–integrin axes",
            "GRN-regression virtual knockdown prioritizes candidate perturbation-sensitive regulators",
            "Cross-species spatial analysis supports DA-FAP localization in dystrophic and injured niches",
            "Multi-cohort human bulk validation supports reproducible DA-FAP enrichment",
            "Sensitivity analyses support robustness while flagging dataset-specific caveats",
        }:
            doc.add_paragraph(text, style="Heading 2")
        elif text.startswith(("Data availability.", "Competing interests.", "Funding.")):
            add_paragraph(doc, text)
        else:
            add_paragraph(doc, text)

    add_paragraph(
        doc,
        "Ethics approval and consent to participate. Not required for this secondary analysis of publicly available de-identified transcriptomic datasets.",
    )
    add_paragraph(
        doc,
        "Author contributions. M.C. conceived the study, designed and implemented the computational analysis, interpreted the results, and wrote the manuscript.",
    )
    add_paragraph(
        doc,
        "Code availability. The reproducible analysis scripts and release packages are organized in the accompanying project directory.",
    )

    doc.add_paragraph("References", style="Heading 1")
    for reference in parts["references"]:
        add_paragraph(doc, reference)

    doc.add_paragraph("Tables", style="Heading 1")
    table_widths = [
        [1500, 1100, 2200, 1400, 1300, 1860],
        [650, 1050, 950, 1100, 1200, 1600, 2810],
        [2000, 950, 1150, 1850, 1300, 1100],
    ]
    for i, caption in enumerate(parts["table_captions"]):
        if i > 0:
            doc.add_page_break()
        add_paragraph(doc, caption)
        if i < len(source.tables):
            copy_source_table(doc, source.tables[i], table_widths[i])

    doc.add_paragraph("Figure legends", style="Heading 1")
    for legend in parts["figure_legends"]:
        add_paragraph(doc, legend)

    doc.core_properties.title = TITLE
    doc.core_properties.author = "Author Name"
    doc.save(SUBMISSION_DOCX)


def add_markdown_as_docx(doc: Document, markdown_path: Path) -> None:
    if not markdown_path.exists():
        return
    for raw_line in markdown_path.read_text(encoding="utf-8").splitlines():
        line = clean_text(raw_line.strip()).replace("`", "")
        if not line:
            continue
        if line.lower() in {"supplementary methods", "# supplementary methods"}:
            continue
        if line.startswith("# "):
            doc.add_paragraph(line[2:].strip(), style="Heading 1")
        elif line.startswith("## "):
            doc.add_paragraph(line[3:].strip(), style="Heading 2")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:].strip(), style="Heading 3")
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(line[2:].strip())
        else:
            add_paragraph(doc, line)


def compact_value(value) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, float):
        if abs(value) >= 1000 or abs(value) < 0.001 and value != 0:
            return f"{value:.3g}"
        return f"{value:.4f}".rstrip("0").rstrip(".")
    text = str(value)
    if len(text) > 220:
        return text[:217] + "..."
    return text


def add_dataframe_table(doc: Document, df: pd.DataFrame, max_rows: int | None = None) -> None:
    display = df if max_rows is None else df.head(max_rows)
    rows, cols = display.shape
    if rows == 0 or cols == 0:
        add_paragraph(doc, "No rows available.")
        return

    table = doc.add_table(rows=rows + 1, cols=cols)
    table.style = "Table Grid"
    for j, col in enumerate(display.columns):
        table.cell(0, j).text = str(col)
    for i, (_, row) in enumerate(display.iterrows(), start=1):
        for j, col in enumerate(display.columns):
            table.cell(i, j).text = compact_value(row[col])

    base_width = max(520, int(9360 / max(cols, 1)))
    widths = [base_width] * cols
    if cols >= 10:
        font_size = 5.5
    elif cols >= 7:
        font_size = 6.5
    else:
        font_size = 8
    set_table_geometry(table, widths, font_size=font_size)


def preview_columns(df: pd.DataFrame) -> list[str]:
    priority = [
        "dataset_id",
        "repository",
        "species",
        "tissue",
        "condition",
        "disease_or_condition",
        "data_type",
        "inclusion_decision",
        "module",
        "regulator",
        "fap_subtype",
        "gene",
        "sender",
        "receiver",
        "ligand",
        "receptor",
        "pathway",
        "contrast_label",
        "random_effect_hedges_g",
        "pvalue",
    ]
    chosen = [col for col in priority if col in df.columns]
    if len(chosen) < 6:
        for col in df.columns:
            if col not in chosen:
                chosen.append(col)
            if len(chosen) >= 8:
                break
    return chosen[:8]


def build_supplementary_workbook(table_paths: list[Path]) -> None:
    with pd.ExcelWriter(SUPPLEMENT_XLSX, engine="openpyxl") as writer:
        for path in table_paths:
            number_match = re.search(r"supplementary_table_(\d+)_", path.name)
            number = int(number_match.group(1)) if number_match else 0
            sheet_name = f"Table {number}"
            df = pd.read_csv(path)
            df.to_excel(writer, index=False, sheet_name=sheet_name)


def build_supplementary_materials() -> None:
    doc = Document()
    apply_base_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Supplementary materials")
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    title_run.font.size = Pt(16)

    add_paragraph(doc, TITLE)
    add_paragraph(doc, "Author Name")
    add_paragraph(doc, "Correspondence: author@example.com")

    doc.add_paragraph("Supplementary methods", style="Heading 1")
    add_markdown_as_docx(doc, PROJECT_ROOT / "manuscript" / "supplementary_methods.md")

    doc.add_paragraph("Supplementary table legends and source-data index", style="Heading 1")
    add_paragraph(
        doc,
        "The complete machine-readable supplementary tables are provided in the accompanying Excel workbook "
        "'supplementary source tables.xlsx' and in the project CSV source tables. Concise tables are embedded "
        "below; large source-data tables are summarized in the document to preserve readability.",
    )

    table_paths = sorted(
        (PROJECT_ROOT / "results" / "tables").glob("supplementary_table_*.csv"),
        key=lambda p: int(re.search(r"supplementary_table_(\d+)_", p.name).group(1)),
    )
    build_supplementary_workbook(table_paths)

    overview_rows = []
    for path in table_paths:
        number = int(re.search(r"supplementary_table_(\d+)_", path.name).group(1))
        df = pd.read_csv(path)
        overview_rows.append(
            {
                "Supplementary item": f"Supplementary Table {number}",
                "Description": SUPPLEMENTARY_TABLE_CAPTIONS.get(number, ""),
                "Rows": df.shape[0],
                "Columns": df.shape[1],
                "Source file": path.name,
            }
        )
    add_dataframe_table(doc, pd.DataFrame(overview_rows))

    for path in table_paths:
        number = int(re.search(r"supplementary_table_(\d+)_", path.name).group(1))
        df = pd.read_csv(path)
        doc.add_page_break()
        doc.add_paragraph(f"Supplementary Table {number}", style="Heading 1")
        add_paragraph(doc, SUPPLEMENTARY_TABLE_CAPTIONS.get(number, ""))
        add_paragraph(doc, f"Source file: {path.name}. Dimensions: {df.shape[0]} rows x {df.shape[1]} columns.")
        if df.shape[0] <= 120 and df.shape[1] <= 10:
            add_dataframe_table(doc, df)
        else:
            columns = pd.DataFrame({"Column": list(df.columns)})
            add_paragraph(
                doc,
                "Large source-data table. Column names are listed below; full rows are provided in the accompanying Excel workbook.",
            )
            add_dataframe_table(doc, columns)
            preview_rows = min(15, df.shape[0])
            selected_columns = preview_columns(df)
            add_paragraph(doc, f"Preview of the first {preview_rows} rows using selected readable columns:")
            add_dataframe_table(doc, df.loc[:, selected_columns].head(preview_rows))

    doc.core_properties.title = "Supplementary materials"
    doc.core_properties.author = "Author Name"
    doc.save(SUPPLEMENT_DOCX)


def main() -> int:
    if not SOURCE_DOCX.exists():
        print(f"Source manuscript not found: {SOURCE_DOCX}", file=sys.stderr)
        return 1
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_submission_manuscript()
    build_supplementary_materials()
    print(f"Wrote {SUBMISSION_DOCX}")
    print(f"Wrote {SUPPLEMENT_DOCX}")
    print(f"Wrote {SUPPLEMENT_XLSX}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

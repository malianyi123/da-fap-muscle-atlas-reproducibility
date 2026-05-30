#!/usr/bin/env python3
"""Build a clean upgraded submission package after reviewer-recommended analyses."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "投稿_升级版"
SOURCE_MAIN = ROOT / "manuscript" / "manuscript_submission_revised.docx"
TABLE_DIR = ROOT / "results" / "tables"
FIG_DIR = ROOT / "results" / "figures"

TITLE = "Perturbation-guided transcriptomic atlases prioritize degeneration-associated fibro-adipogenic progenitor programs in skeletal muscle"


SUPP_TABLE_CAPTIONS = {
    1: "Public dataset inventory, metadata fields, inclusion decisions, and rationale.",
    2: "Single-cell and single-nucleus quality-control thresholds and retained cells.",
    3: "Major cell-type marker-module score summaries used for atlas annotation.",
    4: "FAP subtype marker genes ranked by differential expression statistics.",
    5: "DA-FAP and component gene modules used for scoring.",
    6: "Curated ligand-receptor interaction scores across sender and receiver groups.",
    7: "Candidate regulator expression and correlation with DA-FAP component scores.",
    8: "GRN-regression virtual knockdown ranking across candidate regulators.",
    9: "Spatial co-localization statistics in GSE225766 Visium samples.",
    10: "Bulk validation effect sizes across individual cohorts and contrasts.",
    11: "Subtype-level FAP module programs.",
    12: "Leave-one-dataset-out DA-FAP label-recovery sensitivity analysis.",
    13: "DA-FAP module-gene sensitivity and bootstrap stability analysis.",
    14: "Correlations between DA-FAP score and quality-control or artifact-associated covariates.",
    15: "Integration-method sensitivity comparing uncorrected PCA and Harmony representations.",
    16: "Random-effects meta-analysis of module scores across primary human bulk cohorts.",
    17: "Null-label permutation analysis for DA-FAP regulator specificity in the GRN-regression perturbation screen.",
    18: "Matched-expression regulator-control analysis for perturbation-sensitivity magnitude.",
    19: "Bootstrap/subsampling stability of regulator ranking, including top-5/top-10 overlap and rank correlation summaries.",
    20: "Donor/sample-level paired re-analysis of pseudotime and DA-FAP score comparisons where metadata permitted.",
}

SUPP_FIGURES = [
    (
        1,
        "phase3_umap_by_leiden.png",
        "Unsupervised clustering of the integrated atlas. UMAP of the 121,520-cell integrated human skeletal muscle atlas coloured by Leiden cluster.",
    ),
    (
        2,
        "phase7_regulator_target_proxy_network.png",
        "Correlation-based candidate regulator network preceding GRN-regression virtual knockdown.",
    ),
    (
        3,
        "phase8_virtual_ko_proxy_heatmap.png",
        "Correlation-based virtual knockout proxy scores for candidate regulators across DA-FAP component programs.",
    ),
    (
        4,
        "phase8b_grn_virtual_knockdown_ranking.png",
        "GRN-regression virtual knockdown ranking by predicted DA-FAP program attenuation.",
    ),
    (
        5,
        "phase9_spatial_gse225766_da_colocalization.png",
        "Spatial co-localization of the DA-FAP signature with fibrotic/ECM, inflammatory-remodeling, FAP-stromal, macrophage, adipogenic, and myofiber signatures.",
    ),
    (
        6,
        "phase9_spatial_gse225766_da_score_by_condition.png",
        "Spot-level DA-FAP score distributions across wild-type, mdx dystrophic, and post-injury muscle samples in GSE225766.",
    ),
    (
        7,
        "phase10_bulk_gse164471_da_score_age_correlation.png",
        "Association between the bulk DA-FAP composite score and donor age in the independent human aging cohort GSE164471.",
    ),
    (
        8,
        "phase10_bulk_gse164471_module_heatmap.png",
        "Module-score heatmap in the GSE164471 cohort.",
    ),
    (
        9,
        "suppfig9_grn_robustness.png",
        "Robustness of the GRN-regression regulator screen, including null-label permutation, matched-expression controls, and bootstrap ranking stability.",
    ),
    (
        10,
        "suppfig10_donor_level.png",
        "Donor/sample-level re-analysis of key single-cell comparisons where metadata permitted.",
    ),
]


def set_font(run, size=11, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size in [("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)


def add_para(doc: Document, text: str, style: str | None = None, bold=False):
    p = doc.add_paragraph(style=style)
    r = p.add_run(str(text))
    set_font(r, bold=bold)
    return p


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], font_size=8.0) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    for row_idx, row in enumerate(table.rows):
        if row_idx == 0:
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:tblHeader")) is None:
                tr_pr.append(OxmlElement("w:tblHeader"))
        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths_dxa[min(col_idx, len(widths_dxa) - 1)])
            if row_idx == 0:
                set_cell_shading(cell, "EDEDED")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    set_font(r, size=font_size, bold=(row_idx == 0))


def clean_value(value) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, float):
        if abs(value) >= 1000 or (abs(value) < 0.001 and value != 0):
            return f"{value:.3g}"
        return f"{value:.4f}".rstrip("0").rstrip(".")
    text = str(value)
    return text[:217] + "..." if len(text) > 220 else text


def add_dataframe(doc: Document, df: pd.DataFrame) -> None:
    rows, cols = df.shape
    table = doc.add_table(rows=rows + 1, cols=cols)
    table.style = "Table Grid"
    for j, col in enumerate(df.columns):
        table.cell(0, j).text = str(col)
    for i, (_, row) in enumerate(df.iterrows(), start=1):
        for j, col in enumerate(df.columns):
            table.cell(i, j).text = clean_value(row[col])
    widths = [max(560, int(9360 / max(cols, 1)))] * cols
    size = 5.8 if cols >= 9 else 7.0 if cols >= 6 else 8.0
    set_table_geometry(table, widths, font_size=size)


def preview_columns(df: pd.DataFrame) -> list[str]:
    preferred = [
        "record_type",
        "comparison",
        "regulator",
        "module",
        "dataset_id",
        "donor",
        "fap_subtype",
        "gene",
        "sender",
        "receiver",
        "ligand",
        "receptor",
        "pathway",
        "observed_sensitivity",
        "sensitivity",
        "perm_p",
        "perm_fdr",
        "top5_selection_freq",
        "top10_selection_freq",
        "wilcoxon_p",
        "pvalue",
        "interpretation",
    ]
    cols = [c for c in preferred if c in df.columns]
    for col in df.columns:
        if col not in cols:
            cols.append(col)
        if len(cols) >= 8:
            break
    return cols[:8]


def table_paths() -> list[Path]:
    paths = sorted(
        TABLE_DIR.glob("supplementary_table_*.csv"),
        key=lambda p: int(re.search(r"supplementary_table_(\d+)_", p.name).group(1)),
    )
    return [p for p in paths if int(re.search(r"supplementary_table_(\d+)_", p.name).group(1)) <= 20]


def build_workbook(paths: list[Path]) -> None:
    with pd.ExcelWriter(OUT / "Supplementary Source Tables.xlsx", engine="openpyxl") as writer:
        for path in paths:
            n = int(re.search(r"supplementary_table_(\d+)_", path.name).group(1))
            pd.read_csv(path).to_excel(writer, sheet_name=f"Table {n}", index=False)


def build_main_manuscript() -> None:
    doc = Document(SOURCE_MAIN)
    style_document(doc)
    if doc.paragraphs:
        doc.paragraphs[0].text = TITLE
    replacements = {
        "Figure 1. Study design and public dataset map. Public skeletal muscle transcriptomic resources were organized into single-cell or single-nucleus discovery datasets, spatial transcriptomics support datasets, and bulk validation datasets. The analysis framework proceeds from public atlas integration to FAP-like stromal subclustering, DA-FAP score definition, trajectory and communication analysis, GRN-regression virtual knockdown, spatial signature support, bulk meta-validation, and sensitivity analysis.": (
            "Figure 1. Study design and evidence architecture. Public single-cell/single-nucleus, spatial, and bulk skeletal muscle transcriptomic resources were organized into a staged evidence chain: atlas integration, FAP-state definition, state dynamics, GRN-regression perturbation, spatial and bulk validation, and a dedicated robustness layer comprising null-label permutation, matched-expression controls, bootstrap regulator-ranking stability, donor-level aggregation, and integration-method sensitivity."
        ),
        "Figure 5. Regulatory network and GRN-regression virtual knockdown. (A) Directed GRN-regression network for prioritized perturbation-sensitive regulators. (B) Predicted module-score changes after simulated regulator knockdown to fifth-percentile expression in DA-FAP cells. (C) Candidate ranking by predicted DA-FAP attenuation; negative values indicate predicted module attenuation. Results are computational predictions and do not establish experimental causality.": (
            "Figure 5. GRN-regression perturbation screen with negative controls. (A) Perturbation framework for fitting target-gene GRN regressors and simulating transcription-factor knockdown in DA-FAP cells. (B) Candidate ranking by predicted DA-FAP attenuation, with null-label-permutation-supported regulators highlighted. (C) Program-level predicted responses for adipogenic, fibrotic/ECM, and inflammatory-remodeling modules. (D) Null-label permutation comparing observed regulator sensitivity with label-permuted distributions. (E) Matched-expression control analysis comparing candidate regulators with expression-matched non-priority genes. (F) Stable prioritized regulator-target nodes summarized with bootstrap and model-concordance metrics. Results are computational predictions and do not establish experimental causality."
        ),
        "p = 0.008; Supplementary Figure 10B": "p = 0.0066; Supplementary Figure 10B",
        "The bulk evidence thus supports a reproducible, transcriptome-level DA-FAP program across independent aging and sarcopenia cohorts.": (
            "The bulk evidence thus supports a reproducible, transcriptome-level DA-FAP program across independent aging and sarcopenia cohorts. Together with the single-cell, spatial, communication, and GRN-perturbation analyses, these results support the working DA-FAP degeneration model summarized in Figure 8."
        ),
        "This framework is directly relevant to MRI-defined muscle degeneration, particularly paraspinal fatty infiltration and fibrosis, which are strongly associated with chronic low back pain and spinal degeneration but are usually inferred from imaging without molecular resolution [2];": (
            "This framework, summarized in Figure 8, is directly relevant to MRI-defined muscle degeneration, particularly paraspinal fatty infiltration and fibrosis, which are strongly associated with chronic low back pain and spinal degeneration but are usually inferred from imaging without molecular resolution [2];"
        ),
        "Sensitivity analyses tested the stability of the DA-FAP signature (Figure 9).": (
            "Sensitivity analyses tested the stability of the DA-FAP signature."
        ),
    }
    for p in doc.paragraphs:
        for old, new in replacements.items():
            if old in p.text:
                p.text = p.text.replace(old, new)
    doc.core_properties.title = TITLE
    doc.save(OUT / "Manuscript.docx")


def build_supplementary_methods_tables(paths: list[Path]) -> None:
    doc = Document()
    style_document(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Supplementary Methods and Tables")
    set_font(r, 16, True)
    add_para(doc, TITLE)
    add_para(doc, "Author Name")
    add_para(doc, "Correspondence: author@example.com")

    doc.add_paragraph("Supplementary methods", style="Heading 1")
    method_sections = [
        (
            "Project setup and data processing",
            "The project used a Python-first reproducible analysis framework centered on Scanpy, AnnData, harmonypy, pandas, scipy, scikit-learn, matplotlib, and seaborn. Public single-cell/single-nucleus, spatial transcriptomic, and bulk transcriptomic datasets were standardized into harmonized metadata fields and processed objects. All primary analyses read paths from config/config.yaml and wrote intermediate objects, tables, logs, and figures to the project results directory.",
        ),
        (
            "Single-cell integration and FAP-state analysis",
            "Human cells and nuclei were filtered using dataset-appropriate quality-control thresholds, normalized, log-transformed, integrated with Harmony, clustered, and annotated using marker-module scores and reference labels where available. FAP-like cells were subset and re-integrated, and DA-FAP scores were defined as the mean of z-scored adipogenic, fibrotic/ECM, and inflammatory-remodeling modules. DA-FAPs were interpreted as a degeneration-associated FAP-like state rather than a newly defined cell type.",
        ),
        (
            "GRN-regression virtual knockdown",
            "For each DA-FAP marker or module target, an ExtraTrees regression model was trained from FAP-like expression using candidate regulators as predictors, excluding target self-predictors. Regulators were simulated as knocked down to fifth-percentile expression in DA-FAP cells, and downstream target/module changes were predicted from the fitted models. These predictions nominate candidate perturbation-sensitive nodes and do not constitute experimental knockout evidence.",
        ),
        (
            "Null-label permutation and matched-expression controls",
            "To test whether predicted regulator effects were DA-FAP-state-specific, DA-FAP labels were permuted while preserving the number of DA-FAP cells, and observed regulator sensitivity was compared with label-permuted null distributions. Matched-expression controls compared candidate regulators with non-priority genes of similar expression magnitude to assess whether large predicted effects could arise from expression level or model structure rather than specific regulator biology.",
        ),
        (
            "Bootstrap regulator-ranking stability",
            "Regulator-ranking robustness was assessed by bootstrap/subsampling iterations. Stability was summarized by top-5 and top-10 overlap with the observed ranking, Spearman rank correlation, and regulator-level top-5/top-10 selection frequency. A ridge-regression surrogate was also compared with the ExtraTrees ranking to assess cross-model concordance.",
        ),
        (
            "Donor/sample-level aggregation",
            "Where donor or sample metadata permitted, key single-cell comparisons were aggregated to donor/sample-level medians before testing. Pseudotime and DA-FAP score differences between DA-FAP and non-DA-FAP cells were compared across units with at least ten cells per group using Wilcoxon signed-rank tests. These analyses address pseudoreplication and are interpreted alongside cell-level results.",
        ),
        (
            "Spatial and bulk validation",
            "Mouse Visium data were processed as cross-species spatial support rather than human spatial validation. DA-FAP and partner signatures were scored at the spot level and tested for condition differences and co-localization. Independent human bulk cohorts were scored for DA-FAP and component modules, and primary aging/sarcopenia contrasts were combined using random-effects meta-analysis.",
        ),
    ]
    for heading, text in method_sections:
        doc.add_paragraph(heading, style="Heading 2")
        add_para(doc, text)

    doc.add_paragraph("Supplementary table index", style="Heading 1")
    rows = []
    for path in paths:
        n = int(re.search(r"supplementary_table_(\d+)_", path.name).group(1))
        df = pd.read_csv(path)
        rows.append({"Item": f"Supplementary Table {n}", "Description": SUPP_TABLE_CAPTIONS[n], "Rows": df.shape[0], "Columns": df.shape[1], "Source file": path.name})
    add_dataframe(doc, pd.DataFrame(rows))

    for path in paths:
        n = int(re.search(r"supplementary_table_(\d+)_", path.name).group(1))
        df = pd.read_csv(path)
        doc.add_page_break()
        doc.add_paragraph(f"Supplementary Table {n}", style="Heading 1")
        add_para(doc, SUPP_TABLE_CAPTIONS[n])
        add_para(doc, f"Source file: {path.name}. Dimensions: {df.shape[0]} rows x {df.shape[1]} columns.")
        if df.shape[0] <= 80 and df.shape[1] <= 9:
            add_dataframe(doc, df)
        else:
            add_para(doc, "Large source-data table. Column names and a readable preview are shown below; complete rows are provided in Supplementary Source Tables.xlsx.")
            add_dataframe(doc, pd.DataFrame({"Column": list(df.columns)}))
            add_para(doc, f"Preview of first {min(15, df.shape[0])} rows using selected columns:")
            add_dataframe(doc, df.loc[:, preview_columns(df)].head(15))

    doc.core_properties.title = "Supplementary Methods and Tables"
    doc.save(OUT / "Supplementary Methods and Tables.docx")


def build_supplementary_figures_doc() -> None:
    doc = Document()
    style_document(doc)
    doc_image_dir = OUT / "_doc_images"
    doc_image_dir.mkdir(exist_ok=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Supplementary Figures")
    set_font(r, 16, True)
    add_para(doc, TITLE)
    for i, filename, caption in SUPP_FIGURES:
        if i > 1:
            doc.add_page_break()
        doc.add_paragraph(f"Supplementary Figure {i}", style="Heading 1")
        image_path = FIG_DIR / filename
        if image_path.exists():
            flattened = doc_image_dir / f"supplementary_figure_{i}.png"
            flatten_png(image_path, flattened)
            shape = doc.add_picture(str(flattened), width=Inches(6.4))
            doc_pr = shape._inline.docPr
            doc_pr.set("title", f"Supplementary Figure {i}")
            doc_pr.set("descr", caption)
        add_para(doc, f"Supplementary Figure {i}. {caption}")
    doc.core_properties.title = "Supplementary Figures"
    doc.save(OUT / "Supplementary Figures.docx")


def flatten_png(source: Path, destination: Path) -> None:
    image = Image.open(source)
    if image.mode in {"RGBA", "LA"}:
        background = Image.new("RGB", image.size, "white")
        alpha = image.getchannel("A")
        background.paste(image.convert("RGB"), mask=alpha)
        background.save(destination)
    else:
        image.convert("RGB").save(destination)


def copy_figures() -> None:
    main_dir = OUT / "Figures"
    supp_dir = OUT / "Supplementary Figures"
    main_dir.mkdir(parents=True, exist_ok=True)
    supp_dir.mkdir(parents=True, exist_ok=True)
    for i in range(1, 10):
        for suffix in [".png", ".pdf", ".svg"]:
            matches = sorted(FIG_DIR.glob(f"figure{i}_*{suffix}"))
            if matches:
                shutil.copy2(matches[0], main_dir / f"Figure_{i}{suffix}")
    for i, filename, _ in SUPP_FIGURES:
        src = FIG_DIR / filename
        if src.exists():
            flatten_png(src, supp_dir / f"Supplementary_Figure_{i}.png")
        for suffix in [".pdf", ".svg"]:
            sibling = src.with_suffix(suffix)
            if sibling.exists():
                shutil.copy2(sibling, supp_dir / f"Supplementary_Figure_{i}{suffix}")


def write_readme() -> None:
    text = f"""# Upgraded DA-FAP submission package

This package incorporates the reviewer-recommended pre-submission upgrades:

1. Null-label permutation for DA-FAP labels in the GRN-regression virtual knockdown framework.
2. Matched-expression regulator controls.
3. Bootstrap/subsampling stability of regulator ranking, including top-5/top-10 overlap and Spearman rank correlation.
4. Donor/sample-level aggregation for key single-cell comparisons where metadata permit.
5. Redesigned Figure 1 and Figure 5 for a stronger high-impact journal presentation.

Core files:

- `Manuscript.docx`
- `Supplementary Methods and Tables.docx`
- `Supplementary Figures.docx`
- `Supplementary Source Tables.xlsx`
- `Figures/`
- `Supplementary Figures/`

All GRN perturbation outputs remain computational predictions and should not be interpreted as experimental causality.
"""
    (OUT / "README.md").write_text(text, encoding="utf-8")


def main() -> int:
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)
    paths = table_paths()
    build_workbook(paths)
    build_main_manuscript()
    build_supplementary_methods_tables(paths)
    build_supplementary_figures_doc()
    copy_figures()
    write_readme()
    shutil.rmtree(OUT / "_doc_images", ignore_errors=True)
    print(f"Built upgraded submission package: {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
